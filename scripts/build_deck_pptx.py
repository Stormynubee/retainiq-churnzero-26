"""Generate ChurnZero_RetainIQ_Presentation.pptx from slide defs and deck charts."""

from __future__ import annotations

import json
import os
import sys
from pathlib import Path
from typing import Any

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

from pptx import Presentation
from pptx.util import Inches, Pt

from retainiq import config
from retainiq.deck_metrics import DeckMetrics, load_deck_metrics

CHARTS = ROOT / "deck" / "charts"
DEFAULT_OUT = ROOT / "deck" / "ChurnZero_RetainIQ_Presentation.pptx"

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)


def _slide11_bullets() -> list[str]:
    path = config.DATA_PROCESSED / "slide11_customer_story.json"
    if path.is_file():
        story = json.loads(path.read_text(encoding="utf-8"))
        return [
            f"Customer {story['customer_id']} · segment: {story['segment']}",
            f"P(churn) = {story['churn_proba']} · CATE = +{story['cate']}",
            "High uplift in persuadable band (CATE + risk), not mass blast",
            story["talking_points"][0],
            story["talking_points"][2],
        ]
    return [
        "Run: python -m scripts.pick_persuadable_story",
        "(after build_artifacts)",
        "High uplift + risk band → targeted RM call",
        "Not a mass promotional blast",
    ]


def _slides_spec(m: DeckMetrics) -> list[dict[str, Any]]:
    t_label = f"{m.threshold:.3f}".rstrip("0").rstrip(".")
    if len(t_label.split(".")[-1]) < 3 and m.threshold < 0.01:
        t_label = f"{m.threshold:.3f}"

    return [
        {
            "title": "RetainIQ — Who to call, and what it costs",
            "bullets": [
                "Hansraj Tiwari & swayangjeet nayak",
                "ChurnZero 26 · IIT Kharagpur",
            ],
        },
        {
            "title": "The 80:1 problem",
            "bullets": [
                "Missed churner (FN): ₹40,000",
                "Unnecessary call (FP): ₹500",
                "Cost ratio 80:1 — threshold 0.5 is wrong for this problem",
            ],
        },
        {
            "title": "Dataset snapshot",
            "bullets": [
                "8,101 train · 97 features · 16.1% churn",
                "2,026 test customers · submission 16.6% positive",
                "Inactivity and balance stress separate churners",
            ],
        },
        {
            "title": "Where churn concentrates",
            "bullets": [
                "Higher churn: low tenure, low digital engagement, complaints",
                "Lower churn: stable tenure + active digital footprint",
                "Segments guide playbook priority",
            ],
        },
        {
            "title": "RetainIQ pipeline",
            "bullets": [
                "LGB + CatBoost OOF → meta on OOF → Platt calibration",
                "Rank → churn_probability; calibrated + cost → churn_prediction",
                "IPTW uplift · fairness (gender, region)",
            ],
        },
        {
            "title": "Model scores plateau — cutoff still matters",
            "bullets": [
                f"Logistic → LGB → stack → calibrated: PR-AUC ≈ {m.pr_auc}",
                "AUC gains are flat on this dataset",
                "Where we win: cost-optimal cutoff + uplift + playbook",
            ],
        },
        {
            "title": "Out-of-fold results",
            "table": [
                ["", f"t = {t_label}", "t = 0.5"],
                ["PR-AUC", str(m.pr_auc), str(m.pr_auc)],
                [
                    "Business cost (INR)",
                    f"{m.cost_optimal_inr:,}",
                    f"{m.cost_naive_inr:,}",
                ],
                ["Recall", "99.9%", "99.6%"],
            ],
            "bullets": ["5-fold OOF · seed 42 · primary metric: PR-AUC"],
        },
        {
            "title": "Same model, different cutoff",
            "bullets": [
                f"Cost-optimal threshold {t_label} vs naive 0.5",
                f"₹{m.savings_inr:,} saved on 8,101 customers (~{m.savings_pct:.0f}%)",
            ],
            "image": CHARTS / "08_cost_curve.png",
        },
        {
            "title": "What drives churn",
            "bullets": [
                "Actionable: logins, complaints, campaigns, RM touches",
                "Structural: tenure (monitor only)",
                "Top: total_digital_logins (~30% LGB gain)",
            ],
            "image": CHARTS / "09_feature_importance.png",
        },
        {
            "title": "Who is worth calling?",
            "bullets": [
                f"{m.persuadable_n} persuadables — offer helps (CATE +{m.persuadable_avg_cate:.2f} avg)",
                f"{m.sleeping_dog_n} sleeping-dogs — offer may hurt (CATE {m.sleeping_dog_avg_cate:.2f} avg)",
                "Offers not randomized — IPTW; directional only",
            ],
            "image": CHARTS / "10_uplift_quadrant.png",
        },
        {
            "title": "Case study — persuadable segment",
            "bullets": _slide11_bullets(),
        },
        {
            "title": "Fairness at operating threshold",
            "bullets": [
                f"Gender DP gap: {m.gender_dp_gap_pct}% · Region: {m.region_dp_gap_pct}% (limit 10%)",
                "Equal opportunity gaps ≈ 0",
            ],
            "image": CHARTS / "12_fairness.png",
        },
        {
            "title": "Monday-morning playbook",
            "bullets": [
                "Persuadable → RM call + waiver",
                "Sleeping-dog → no promotional contact",
                "Lost-cause / sure-thing → low-touch or cross-sell",
                f"ROI: ₹{m.savings_inr / 1000:.1f}k saved (~₹{m.savings_inr / 8101:.0f}/customer)",
            ],
        },
        {
            "title": "Production path",
            "bullets": [
                "Weekly CRM scores · quarterly retrain · PSI on drivers",
                "python -m scripts.train · 58 pytest tests + CI",
                "Limits: observational uplift, PR-AUC ~1.0 on this dataset",
            ],
        },
        {
            "title": "RetainIQ — Thank you",
            "bullets": [
                "Questions?",
                "github.com/Stormynubee/retainiq-churnzero-26",
            ],
        },
    ]


def _add_title_content_slide(
    prs: Presentation,
    title: str,
    bullets: list[str],
    image_path: Path | None = None,
    table_rows: list[list[str]] | None = None,
) -> None:
    layout = prs.slide_layouts[1]
    slide = prs.slides.add_slide(layout)
    slide.shapes.title.text = title

    body = slide.placeholders[1]
    tf = body.text_frame
    if not table_rows and not image_path:
        tf.text = bullets[0] if bullets else ""
        for line in bullets[1:5]:
            p = tf.add_paragraph()
            p.text = line
            p.font.size = Pt(18)
    elif bullets and not table_rows:
        tf.text = bullets[0]
        for line in bullets[1 : 3 if image_path else 5]:
            p = tf.add_paragraph()
            p.text = line
            p.font.size = Pt(16)
    else:
        tf.text = bullets[0] if bullets else ""
        for line in bullets[1:2]:
            p = tf.add_paragraph()
            p.text = line
            p.font.size = Pt(14)

    top = Inches(1.35)
    if table_rows:
        rows, cols = len(table_rows), len(table_rows[0])
        left = Inches(0.8)
        width = Inches(11.5)
        height = Inches(0.35 * rows)
        shape = slide.shapes.add_table(rows, cols, left, top, width, height)
        table = shape.table
        for r, row in enumerate(table_rows):
            for c, val in enumerate(row):
                cell = table.cell(r, c)
                cell.text = val
                if r == 0:
                    for paragraph in cell.text_frame.paragraphs:
                        paragraph.font.bold = True
                        paragraph.font.size = Pt(14)
                else:
                    for paragraph in cell.text_frame.paragraphs:
                        paragraph.font.size = Pt(14)
        top = top + height + Inches(0.15)

    if image_path and image_path.is_file():
        img_top = top if table_rows else Inches(1.5)
        slide.shapes.add_picture(
            str(image_path),
            Inches(0.6),
            img_top,
            width=Inches(12.0),
        )


def build_deck(out_path: Path | None = None) -> Path:
    m = load_deck_metrics()
    specs = _slides_spec(m)
    out = Path(out_path or os.environ.get("RETAINIQ_DECK_OUT", DEFAULT_OUT))
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H

    for spec in specs:
        _add_title_content_slide(
            prs,
            spec["title"],
            spec.get("bullets", []),
            image_path=spec.get("image"),
            table_rows=spec.get("table"),
        )

    out.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(out))
    return out


def main() -> None:
    path = build_deck()
    print(f"Wrote {path} ({len(_slides_spec(load_deck_metrics()))} slides)")


if __name__ == "__main__":
    main()
